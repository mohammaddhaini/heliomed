#!/usr/bin/env python3
"""
parse-gms-excel.py
Extracts product catalog and image mappings from GMS_PRODUCTS.xlsx
and HELIO-MED-products-batches-01-to-09-COMPLETE/images.

Outputs clean JSON ready for Firebase Firestore and Storage upload pipelines.
"""

import os
import sys
import re
import json
import argparse
import openpyxl

DEFAULT_EXCEL_PATH = r"C:\Users\dr laptop\Downloads\GMS_PRODUCTS.xlsx"
DEFAULT_IMAGES_DIR = r"C:\Users\dr laptop\Downloads\HELIO-MED-products-batches-01-to-09-COMPLETE\images"
DEFAULT_FALLBACK_IMAGES_DIR = r"C:\Users\dr laptop\Downloads\GMS_PRODUCTS\GMS_PRODUCTS\images"
DEFAULT_OUTPUT_JSON = os.path.join(os.path.dirname(__file__), "gms-products-extracted.json")

EXCLUDE_PRICE_TITLES = {
    "rom elbow brace",
    "u gel seat cushion",
    "support stocking calypso (closed toe)"
}

CATEGORY_PRIORITY = [
    ("Pediatric", "Pediatric"),
    ("Pillows", "Pillows"),
    ("Disk Trac", "Disk Trac"),
    ("Vascular", "Vascular"),
    ("Foot Care", "Foot Care"),
    ("Home care", "Home Care"),
    ("Cervical Support", "Cervical Support"),
    ("Wrist&Hand Braces", "Wrist and Hand Braces"),
    ("Shoulder & Elbow", "Shoulder and Elbow"),
    ("Ankle Support", "Ankle Support"),
    ("Knee Support", "Knee Support"),
    ("Back Support", "Back Support"),
    ("Orthopedic", "Orthopedic"),
]

def slug(value: str) -> str:
    s = str(value or "product").strip().lower()
    s = s.replace("&", " and ")
    s = re.sub(r"[^a-z0-9]+", "-", s)
    s = re.sub(r"^-+|-+$", "", s)
    return s[:120] or "product"

def detect_brand(title: str, description: str = "") -> str:
    text = f"{title} {description or ''}".lower()
    if "omron" in text:
        return "OMRON"
    if "rossmax" in text:
        return "Rossmax"
    if "aspen" in text or "vista" in text:
        return "Aspen"
    if "philadelphia" in text:
        return "Philadelphia"
    if "posturex" in text:
        return "Posturex"
    if "vertebradyn" in text:
        return "Vertebradyn"
    if "manu cast" in text or "manu_cast" in text:
        return "Manu Cast"
    if "myolito" in text:
        return "Myolito"
    if "disk-trac" in text or "disk trac" in text:
        return "Disk-Trac"
    return "GMS Medical"

def pick_category(categories_str: str) -> str:
    if not categories_str:
        return "Orthopedic"
    parts = [c.strip() for c in str(categories_str).split(",")]
    for key, target in CATEGORY_PRIORITY:
        if key in parts:
            return target
    return "Orthopedic"

def parse_categories_list(categories_str: str, primary: str) -> list[str]:
    clean = []
    if categories_str:
        for c in str(categories_str).split(","):
            c = c.strip()
            if not c or c.lower() == "all":
                continue
            if c == "Shoulder & Elbow":
                c = "Shoulder and Elbow"
            elif c == "Home care":
                c = "Home Care"
            elif c == "Wrist&Hand Braces":
                c = "Wrist and Hand Braces"
            clean.append(c)
    if primary not in clean:
        clean.insert(0, primary)
    # Remove duplicates preserving order
    seen = set()
    res = []
    for item in clean:
        if item not in seen:
            seen.add(item)
            res.append(item)
    return res

def parse_price(title: str, raw_price: str):
    raw_str = str(raw_price or "").strip()
    if title.lower() in EXCLUDE_PRICE_TITLES:
        return "", None, raw_str
    digits = re.sub(r"[^0-9]", "", raw_str)
    if digits:
        val = int(digits)
        return f"${val}.00", val, raw_str
    return "", None, raw_str

def resolve_product_images(img_col_val: str, images_dir: str, fallback_dir: str = None) -> tuple[str, list[str]]:
    """
    Given the first image path from Excel (e.g. .../images/001_Silicone_Knee_Brace/01.jpg),
    locates the matching folder in images_dir (and fallback_dir if missing),
    and returns (folder_name, [list of full paths to image files]).
    """
    if not img_col_val:
        return "", []

    normalized_path = img_col_val.replace("/", "\\")
    parts = normalized_path.split("\\")
    folder_name = parts[-2] if len(parts) >= 2 else ""

    if not folder_name:
        return "", []

    target_dir = os.path.join(images_dir, folder_name)
    if not os.path.exists(target_dir) and fallback_dir and os.path.exists(fallback_dir):
        fallback_target = os.path.join(fallback_dir, folder_name)
        if os.path.exists(fallback_target):
            target_dir = fallback_target

    if not os.path.exists(target_dir):
        return folder_name, []

    valid_exts = (".webp", ".jpg", ".jpeg", ".png", ".jfif")
    files = [
        os.path.join(target_dir, f)
        for f in sorted(os.listdir(target_dir))
        if f.lower().endswith(valid_exts)
    ]
    return folder_name, files

def extract_products(excel_path: str, images_dir: str, fallback_dir: str = None) -> list[dict]:
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found at: {excel_path}")
    if not os.path.exists(images_dir):
        raise FileNotFoundError(f"Images directory not found at: {images_dir}")

    print(f"Loading Excel workbook: {excel_path} ...")
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    headers = [sheet.cell(row=1, column=col).value for col in range(1, sheet.max_column + 1)]
    header_indices = {h: i + 1 for i, h in enumerate(headers) if h}

    products = []
    total_images_found = 0

    for r in range(2, sheet.max_row + 1):
        def get_val(name):
            col = header_indices.get(name)
            return sheet.cell(row=r, column=col).value if col else None

        title = str(get_val("Product Name") or "").strip()
        if not title:
            continue

        raw_price = get_val("Price")
        new_price, new_price_val, raw_price_str = parse_price(title, raw_price)

        raw_cats = get_val("Categories")
        primary_category = pick_category(raw_cats)
        categories_list = parse_categories_list(raw_cats, primary_category)
        category_slugs = [slug(c) for c in categories_list]

        sku = str(get_val("SKU") or "").strip()
        short_desc = str(get_val("Short Description") or "").strip()
        features = str(get_val("Features") or "").strip()
        specs = str(get_val("Specifications") or "").strip()
        full_desc = str(get_val("Full Description") or "").strip()

        desc_parts = [short_desc]
        if full_desc and full_desc != short_desc:
            desc_parts.append(full_desc)
        if features:
            desc_parts.append(f"Features:\n{features}")
        if specs:
            desc_parts.append(f"Specifications:\n{specs}")
        combined_desc = "\n\n".join(p for p in desc_parts if p)

        brand = detect_brand(title, combined_desc)
        product_url = str(get_val("Product URL") or "").strip()
        product_slug = slug(title)

        img1_val = str(get_val("Image 1") or "").strip()
        folder_name, full_image_paths = resolve_product_images(img1_val, images_dir, fallback_dir)

        # Build relative paths for repo usage (e.g. ./gms-images/001_.../01.webp)
        local_rel_images = []
        for p in full_image_paths:
            fname = os.path.basename(p)
            local_rel_images.append(f"./gms-images/{folder_name}/{fname}")

        total_images_found += len(full_image_paths)

        products.append({
            "id": r - 1,
            "docId": product_slug,
            "title": title,
            "brand": brand,
            "category": primary_category,
            "categories": categories_list,
            "categorySlugs": category_slugs,
            "section": "Medical Supplies",
            "badge": "",
            "oldPrice": "",
            "newPrice": new_price,
            "oldPriceValue": None,
            "newPriceValue": new_price_val,
            "rawPrice": raw_price_str,
            "inventory": 10,
            "available": True,
            "imageFolder": folder_name,
            "localImageFile": local_rel_images[0] if local_rel_images else "",
            "localImageFiles": local_rel_images,
            "absoluteImageFiles": full_image_paths,
            "description": combined_desc,
            "usage": "",
            "warnings": "",
            "sku": sku,
            "sourceUrl": product_url
        })

    print(f"Extracted {len(products)} products from Excel.")
    print(f"Mapped {total_images_found} images across {len(products)} products.")
    return products

def main():
    parser = argparse.ArgumentParser(description="Parse GMS Products Excel and Images directory.")
    parser.add_argument("--excel", default=DEFAULT_EXCEL_PATH, help="Path to GMS_PRODUCTS.xlsx")
    parser.add_argument("--images", default=DEFAULT_IMAGES_DIR, help="Path to images directory")
    parser.add_argument("--fallback-images", default=DEFAULT_FALLBACK_IMAGES_DIR, help="Fallback images directory")
    parser.add_argument("--output", default=DEFAULT_OUTPUT_JSON, help="Output JSON path")
    args = parser.parse_args()

    products = extract_products(args.excel, args.images, args.fallback_images)
    
    # Save standard JSON output (compatible with existing uploader format)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(products, f, indent=2, ensure_ascii=False)
    print(f"Saved extracted data to: {args.output}")

    # Summary
    priced = [p for p in products if p["newPriceValue"] is not None]
    unpriced = [p for p in products if p["newPriceValue"] is None]
    with_imgs = [p for p in products if p["absoluteImageFiles"]]
    print("\nExtraction Summary:")
    print(f"  Total Products: {len(products)}")
    print(f"  Priced:         {len(priced)} (${min(p['newPriceValue'] for p in priced)} - ${max(p['newPriceValue'] for p in priced)})")
    print(f"  Unpriced (Range/Custom): {len(unpriced)} ({', '.join(p['title'] for p in unpriced)})")
    print(f"  Products with Images:    {len(with_imgs)} / {len(products)}")
    print(f"  Total Image Files:       {sum(len(p['absoluteImageFiles']) for p in products)}")

if __name__ == "__main__":
    main()
