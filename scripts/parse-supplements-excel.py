#!/usr/bin/env python3
"""
parse-supplements-excel.py

Extracts supplement products from:
  - Excel: C:\\Users\\dr laptop\\Downloads\\supplements website.xlsx
  - Images: C:\\Users\\dr laptop\\Vs code programs\\client-websites\\tabib-clinc\\scripts\\imgs

Outputs normalized JSON:
  - scripts/supplements-products-extracted.json
"""

import os
import re
import json
import argparse
from pathlib import Path
import openpyxl

DEFAULT_EXCEL_PATH = r"C:\Users\dr laptop\Downloads\supplements website.xlsx"
DEFAULT_IMGS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "imgs")
OUTPUT_JSON_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "supplements-products-extracted.json")

# Explicit mapping from Excel raw item name to clean attributes & image file
PRODUCT_MAP = {
    "GINKO OMEGA 30CAPS": {
        "title": "Ginko Omega 30 Capsules",
        "brand": "Ginko Omega",
        "image": "Ginko-Omega.webp",
        "description": "Dietary supplement containing Ginkgo Biloba extract and Omega-3 fatty acids to support cognitive function, concentration, and cardiovascular health.",
        "category": "Supplements",
        "badge": "Popular"
    },
    "NERVA-Q10 30CAPS": {
        "title": "Nerva-Q10 30 Capsules",
        "brand": "Nerva",
        "image": "NervaQ10_24.webp",
        "description": "Coenzyme Q10 and neuro-supportive micronutrients formulated to maintain healthy nerve function and promote cellular energy production.",
        "category": "Supplements",
        "badge": ""
    },
    "OVABESTFORWOMEN 60TAB": {
        "title": "OvaBest for Women 60 Tablets",
        "brand": "OvaBest",
        "image": "OvaBest.webp",
        "description": "Targeted nutritional support for women formulated to help maintain hormonal balance, ovarian wellness, and metabolic vitality.",
        "category": "Supplements",
        "badge": "Top Seller"
    },
    "RINOPANTENA NASAL OINTMENT 10 G": {
        "title": "Rinopantena Nasal Ointment 10g",
        "brand": "Rinopantena",
        "image": "",
        "description": "Specialized nasal ointment designed to lubricate, moisturize, and promote re-epithelialization of dry or irritated nasal mucous membranes.",
        "category": "Supplements",
        "badge": ""
    },
    "REVAGINAL OVULI-OVULES": {
        "title": "Revaginal Ovules",
        "brand": "Revaginal",
        "image": "",
        "description": "Vaginal ovules specifically formulated to support physiological vaginal flora, optimal pH, and mucosal comfort.",
        "category": "Supplements",
        "badge": ""
    },
    "SILVER PLUS 30 CAPSULES": {
        "title": "Silver Plus 30 Capsules",
        "brand": "Silver Plus",
        "image": "Silver Plus.webp",
        "description": "Comprehensive multivitamin and mineral formulation developed to promote daily vitality, immunity, and overall well-being.",
        "category": "Supplements",
        "badge": ""
    },
    "FLAXAVIT-K2 PLUS 30 CAPSULES": {
        "title": "FlaxaVit-K2 Plus 30 Capsules",
        "brand": "FlaxaVit",
        "image": "FlaxaVit K2.webp",
        "description": "Synergistic complex of Vitamin K2 and essential fatty acids to optimize calcium utilization, bone mineral density, and arterial health.",
        "category": "Supplements",
        "badge": "Premium"
    },
    "ACTIZIM WITH MELATONIN": {
        "title": "Actizim with Melatonin",
        "brand": "Actizim",
        "image": "Actizim.webp",
        "description": "Advanced digestive enzyme blend combined with melatonin to facilitate nighttime digestion and encourage restful, restorative sleep.",
        "category": "Supplements",
        "badge": ""
    },
    "NEUROQ10 64Cmg n": {
        "title": "NeuroQ10 640mg",
        "brand": "NeuroQ10",
        "image": "NeuroQ10_640mg.webp",
        "description": "High-strength Coenzyme Q10 (640mg) supplement formulated for enhanced cellular bioenergetics, cardiovascular health, and antioxidant protection.",
        "category": "Supplements",
        "badge": "High Potency"
    },
    "ARTOGENE 30CAPS": {
        "title": "Artogene 30 Capsules",
        "brand": "Artogene",
        "image": "Artogen.webp",
        "description": "Complete joint support formula containing glucosamine, chondroitin, and trace minerals to nurture cartilage and joint mobility.",
        "category": "Supplements",
        "badge": ""
    },
    "ARTOGENE 60CAPS": {
        "title": "Artogene 60 Capsules",
        "brand": "Artogene",
        "image": "Artogen.webp",
        "description": "Extended-supply joint care formula with glucosamine, chondroitin, and essential cofactors for joint flexibility and cartilage strength.",
        "category": "Supplements",
        "badge": "Value Pack"
    },
    "ENER B 30CAPS": {
        "title": "Ener-B 30 Capsules",
        "brand": "Ener-B",
        "image": "Ener-B.webp",
        "description": "High-potency B-complex supplement formulated to support energy metabolism, reduce fatigue, and sustain nervous system vigor.",
        "category": "Supplements",
        "badge": ""
    },
    "FERTIBEST 30TAB": {
        "title": "FertiBest 30 Tablets",
        "brand": "FertiBest",
        "image": "FertiBest.webp",
        "description": "Targeted blend of antioxidants, amino acids, and minerals clinically selected to support reproductive health and fertility parameters.",
        "category": "Supplements",
        "badge": ""
    },
    "HEMOBEST PLUS 30TAB": {
        "title": "HemoBest Plus 30 Tablets",
        "brand": "HemoBest",
        "image": "Hemobest.webp",
        "description": "Gentle, non-constipating iron complex enriched with Vitamin C, B12, and folic acid to support healthy red blood cell and hemoglobin formation.",
        "category": "Supplements",
        "badge": ""
    },
    "PROBIOVIT Q10": {
        "title": "ProbioVit Q10",
        "brand": "ProbioVit",
        "image": "ProbioVit Q10.webp",
        "description": "Multi-strain probiotic culture combined with CoQ10 to fortify digestive microflora, enhance nutrient absorption, and boost cellular vitality.",
        "category": "Supplements",
        "badge": ""
    },
    "PROBIOVIT C": {
        "title": "ProbioVit C",
        "brand": "ProbioVit",
        "image": "ProbioVit C.webp",
        "description": "Targeted probiotic cultures blended with Vitamin C to bolster gastrointestinal integrity and reinforce immune defense mechanisms.",
        "category": "Supplements",
        "badge": ""
    }
}

def slugify(text: str) -> str:
    s = text.lower().strip()
    s = re.sub(r"[^\w\s-]", "", s)
    s = re.sub(r"[\s_]+", "-", s)
    s = re.sub(r"-+", "-", s)
    return s.strip("-")

def parse_supplements(excel_path: str, imgs_dir: str):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    print(f"Reading Excel: {excel_path}")
    print(f"Headers: {headers}")
    print(f"Total rows: {sheet.max_row}")

    products = []

    for row_idx in range(2, sheet.max_row + 1):
        raw_name = sheet.cell(row=row_idx, column=1).value
        raw_price = sheet.cell(row=row_idx, column=2).value

        if not raw_name:
            continue

        raw_name_clean = str(raw_name).strip()
        meta = PRODUCT_MAP.get(raw_name_clean, {})

        title = meta.get("title", raw_name_clean.title())
        brand = meta.get("brand", "Supplements")
        doc_id = slugify(title)
        category = meta.get("category", "Supplements")
        badge = meta.get("badge", "")
        description = meta.get("description", f"Premium {title} dietary supplement.")
        image_filename = meta.get("image", "")

        # Numeric price handling
        try:
            num_price = float(raw_price) if raw_price is not None else 0.0
            price_str = f"${int(num_price)}.00" if num_price.is_integer() else f"${num_price:.2f}"
        except (ValueError, TypeError):
            num_price = None
            price_str = ""

        # Local image check
        local_image_path = ""
        images_list = []
        if image_filename:
            full_img_path = os.path.join(imgs_dir, image_filename)
            if os.path.exists(full_img_path):
                local_image_path = full_img_path
                images_list = [full_img_path]
            else:
                print(f"[WARN] Image referenced but missing: {full_img_path}")

        # Search text tokens
        search_tokens = [
            title,
            brand,
            category,
            "Supplements",
            "Parapharmacy",
            "Vitamins",
            badge,
            raw_name_clean
        ]
        search_text = " ".join(t.lower() for t in search_tokens if t)

        prod_record = {
            "index": len(products) + 1,
            "excelRow": row_idx,
            "rawName": raw_name_clean,
            "docId": doc_id,
            "title": title,
            "brand": brand,
            "category": category,
            "categories": ["Supplements", "Parapharmacy"],
            "categorySlugs": ["supplements", "parapharmacy"],
            "collections": ["supplements", "parapharmacy"],
            "section": "Parapharmacy",
            "badge": badge,
            "newPrice": price_str,
            "newPriceValue": num_price,
            "oldPrice": "",
            "oldPriceValue": None,
            "inventory": 15,
            "available": True,
            "description": description,
            "usage": "As directed on packaging or by your healthcare professional.",
            "warnings": "Do not exceed recommended daily dosage. Keep out of reach of children.",
            "imageFileName": image_filename,
            "localImagePath": local_image_path,
            "localImageFiles": images_list,
            "searchText": search_text
        }
        products.append(prod_record)

    return products

def main():
    parser = argparse.ArgumentParser(description="Parse supplements Excel and images.")
    parser.add_argument("--excel", default=DEFAULT_EXCEL_PATH, help="Path to supplements website.xlsx")
    parser.add_argument("--imgs", default=DEFAULT_IMGS_DIR, help="Path to images directory")
    parser.add_argument("--output", default=OUTPUT_JSON_PATH, help="Output JSON path")
    args = parser.parse_args()

    products = parse_supplements(args.excel, args.imgs)
    print(f"\nExtracted {len(products)} products.")

    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(products, f, indent=2, ensure_ascii=False)

    print(f"Saved extracted data to: {args.output}")

    # Summary
    with_img = sum(1 for p in products if p["localImagePath"])
    without_img = len(products) - with_img
    print(f"  - With image: {with_img}")
    print(f"  - Without image: {without_img}")

if __name__ == "__main__":
    main()
